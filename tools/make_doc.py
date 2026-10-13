# -*- coding: utf-8 -*-
"""make_doc.py —— 生成《设计文档》docx(匿名, 不含个人信息)
输出: D:\\train\\materials\\设计文档.docx
"""
import os, json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

M=r"D:\train\materials"
ev=json.load(open(os.path.join(M,"eval_results.json"),encoding="utf-8"))
log=json.load(open(r"D:\train\model\train_log.json",encoding="utf-8"))
cmap=json.load(open(r"D:\train\algae_dataset\class_map.json",encoding="utf-8"))
classes=cmap["classes"]; cn=ev["cn"]; summ=cmap["summary"]; pca=ev["per_class_acc"]

doc=Document()
# 中文字体
st=doc.styles["Normal"]; st.font.name="宋体"; st.font.size=Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
for i in range(1,4):
    hs=doc.styles[f"Heading {i}"]; hs.font.name="黑体"; hs.font.color.rgb=RGBColor(0x1F,0x3B,0x63)
    hs.element.rPr.rFonts.set(qn("w:eastAsia"),"黑体")

def H(s,lvl=1): doc.add_heading(s,level=lvl)
def P(s,bold=False,size=11,align=None):
    p=doc.add_paragraph(); r=p.add_run(s); r.bold=bold; r.font.size=Pt(size)
    if align: p.alignment=align
    return p
def table(rows,header=True):
    t=doc.add_table(rows=len(rows),cols=len(rows[0])); t.style="Light Grid Accent 1"
    for i,row in enumerate(rows):
        for j,v in enumerate(row):
            c=t.cell(i,j); c.text=str(v)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size=Pt(10)
                    if header and i==0: r.bold=True
    return t

t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
r=t.add_run("海洋微藻智能识别与赤潮风险提示工具\n设计文档"); r.bold=True; r.font.size=Pt(20)
sub=doc.add_paragraph(); sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=sub.add_run("基于 MindSpore(昇思)迁移学习 · 2026 年海峡大学生人工智能创意大赛"); rr.font.size=Pt(11); rr.font.color.rgb=RGBColor(0x2E,0x75,0xB6)

H("一、项目概述",1)
P("赤潮是福建近岸海域频发的生态灾害,部分赤潮藻产生的毒素可经贝类富集危害人体健康,高密度藻华亦会导致养殖鱼类窒息死亡。传统显微镜人工镜检依赖专业人员、耗时长,难以满足现场快速预警的需求。")
P("本工具面向近岸养殖区、监测站点与公众,采用双分支设计:主分支接受现场最容易获取的海面现状照片,结合水色计算机视觉证据与现场特征问卷做贝叶斯融合,快速给出最可能致灾类群、赤潮风险等级与处置建议;次分支接受显微照片,由基于国产深度学习框架 MindSpore(昇思)与 MindCV 的 MobileNetV3-Small 迁移学习模型做实验室级类别确认,验证集准确率达 %.1f%%。"%(ev["overall_val_acc"]*100))

H("二、数据获取与真实性保障",1)
P("数据真实性是本项目的核心约束。全部训练与验证图像均来自可公开溯源的真实数据源,不使用合成或来源不明的图片:")
doc.add_paragraph("GBIF(全球生物多样性信息网络)公开采集影像:每条记录附带 occurrence key(采集记录号)、记录学名、影像 URL、许可协议(license)与权利人(rightsHolder),逐张写入 provenance.csv;",style="List Bullet")
doc.add_paragraph("CNKI 学术图片库期刊/学位论文显微图:标注图注与文献来源,作为甲藻类的补充;",style="List Bullet")
doc.add_paragraph("分类层级选择:对 GBIF 物种级影像稀疏的类群,在属级(该属全部真实采集影像)进行识别,并在知识库中标注代表种,避免“以他种冒充本种”的标签污染。",style="List Bullet")
P("各类别样本量与验证准确率如下表:")
rows=[["类别(中文)","学名/代表种","训练集","验证集","验证准确率"]]
sci={s:cmap["meta"][s]["sci"] for s in classes}
for s in classes:
    rows.append([cn[s],sci[s],summ[s]["train"],summ[s]["val"],f"{pca[s]*100:.1f}%"])
table(rows)
doc.add_picture(os.path.join(M,"chart_counts.png"),width=Inches(5.6))

H("三、分支一(宏观海面):水色CV × 现场问卷 贝叶斯融合",1)
P("设计动机:真实用户(养殖户、巡查人员、公众)在现场最先获得的是海面现状照片,而非显微照片;但仅凭一张海面照片难以可靠定种。参考赤潮监测规程中综合水色与现场征象的做法,本分支将图像与问卷作为两路独立证据做贝叶斯融合,而非强行端到端分类。")
P("图像证据(无需训练):对海面照去除天空后统计水华色块占比——粉红/砖红(hue≥330 或 <20 且饱和度>0.30)、褐黄/酱色、翠绿/墨绿、乳白/泡沫,按占比阈值判定水色类别并映射为该色系的似然向量。色块存在性检测对岸边局部水华与开阔水华两种构图均稳健,且不依赖宏观标注数据,规避了宏观赤潮图像稀缺的问题。")
P("问卷证据:6 题现场特征——海水颜色、夜间或搅动是否发蓝绿光、是否见死鱼/死虾蟹/贝类开口异常、水面漂浮物、现场气味、近期环境背景(内湾养殖区高温晴朗/大雨后河口/开阔近岸降温大风)。")
P("融合计算:假设集 H={夜光藻型、有毒甲藻型、硅藻型、绿潮、无赤潮}。每个问卷选项与每个水色类别对 H 显式声明一组似然权重(可审计、可写入文档);后验 ∝ 均匀先验 × 图像似然^w × Π(问卷似然),归一化后取最大后验假设,输出致灾类群、风险等级(高/中—高/中/低—中/无)、风险说明与处置建议(是否上报、是否禁捕贝类、是否增氧等)。")
P("案例:一张岸拍粉红色水华照片,水色CV检出粉色块占比约 7% 判定『粉红/砖红』;仅图像证据融合即得夜光藻型 75.7%、有毒甲藻型 18.9%、硅藻型 4.7%,输出风险『中—高』并建议上报与养殖区夜间增氧;叠加『夜间发光=是』等问卷证据后置信度进一步提升。")
P("四级预警分级:为把研判结果直接转化为可执行的行动指令,系统在融合后验之上再叠加一层分级预警。预警指数 = round(后验置信度 × 类群固有风险权重 × 100),其中风险权重把命中类群的危险程度量化——产毒甲藻型=1.0、夜光藻型=0.8、硅藻型=0.6、绿潮=0.4、无赤潮=0;再按指数分档为红(Ⅰ级,≥70)、橙(Ⅱ级,≥50)、黄(Ⅲ级,≥30)、蓝(Ⅳ级,<30)四级预警,每级给出明确的触发条件、含义与响应措施(是否立即上报、是否暂停贝类采捕送检、监测频次、是否增氧或转移网箱、人员防护等)。风险权重表与四级阈值均显式声明,并写入《藻种知识库.xlsx·预警分级表》,与融合似然表一样可审计、可维护。以前述粉红色水华照片为例,仅图像证据即得夜光藻型 75.7%、风险权重 0.8,预警指数 61 → 橙色Ⅱ级预警,提示 48 小时内复核、暂停该点位贝类采捕送检并每日监测。显微分支同样复用这一分级逻辑:识别出藻种后,按知识库中的固有风险等级与识别置信度给出同级预警信号,使宏观与显微两条分支输出统一、可比的预警结论。")

H("四、分支二(显微):模型与方法",1)
P("骨干网络:MobileNetV3-Small(ImageNet 预训练权重),参数量小、推理快,适合教学与端侧部署。")
P("迁移学习策略:冻结骨干全部特征层(206 个参数张量),仅训练分类头(4 个参数张量),在小样本条件下快速收敛并抑制过拟合。")
P("输入与增强:统一中心裁剪并缩放至 224×224 RGB;训练期施加随机裁剪与水平翻转;按 ImageNet 均值/方差归一化。")
P("训练配置:MindSpore PYNATIVE 模式、CPU;损失 CrossEntropyLoss;优化器 Adam(lr=1e-3);batch=16;30 epochs;按验证集准确率自动保存最佳权重。")
rows=[["项目","配置"],["框架","MindSpore 2.x + MindCV"],["骨干","MobileNetV3-Small(预训练)"],
      ["可训练参数","仅分类头"],["输入","224×224×3"],["优化器/学习率","Adam / 1e-3"],
      ["批大小/轮数","16 / 30"],["最佳验证准确率",f"{log['best_val_acc']*100:.1f}%"]]
table(rows)
doc.add_picture(os.path.join(M,"chart_acc_curve.png"),width=Inches(5.6))

H("五、系统实现",1)
P("系统由数据、训练、推理、融合、前端五个模块组成,代码位于 tools/ 目录:")
doc.add_paragraph("prepare_data.py:内容去重→中心裁剪→224×224→按 8:2 划分 train/val→生成 class_map.json;",style="List Bullet")
doc.add_paragraph("train.py:MindSpore 迁移学习训练,自动保存 best.ckpt 与 train_log.json;",style="List Bullet")
doc.add_paragraph("predict.py:加载权重对单张/批量图像推理,并从《藻种知识库.xlsx》联动风险信息;",style="List Bullet")
doc.add_paragraph("fusion.py:宏观分支融合引擎,含水色CV色块检测、6题问卷似然表与贝叶斯融合,纯 NumPy 实现、无需训练;",style="List Bullet")
doc.add_paragraph("app.py:Streamlit 双 Tab 前端——宏观海面(上传+问卷+融合卡+后验证据链)/ 显微(识别+风险卡);低饱和度图像自动提示疑似显微照;",style="List Bullet")
doc.add_paragraph("藻种知识库.xlsx:7 类的赤潮风险等级、生态特性、对养殖影响、毒素类型、爆发规律、监测优先级。",style="List Bullet")
P("前端识别结果与知识库联动:识别出类别后,自动展示该类的风险等级(高/中—高/中/低—中/低)、毒素类型、生态特性、对养殖影响、爆发规律与监测优先级,使“识别”直接转化为“风险提示”。")

H("六、实验结果",1)
_cm=ev['confusion']; _tot=sum(sum(r) for r in _cm); _cor=sum(_cm[i][i] for i in range(len(_cm))); _nc=len(_cm)
P(f"在 {_tot} 张验证图像上,总体准确率为 {ev['overall_val_acc']*100:.2f}%({_cor}/{_tot}),全部 {_nc} 个类别均超过 80% 验收线;最佳单轮验证准确率 {log['best_val_acc']*100:.1f}%。")
doc.add_picture(os.path.join(M,"chart_perclass.png"),width=Inches(5.6))
P("混淆矩阵显示对角线占绝对主导,类间混淆极少;个别误判集中于低清或藻体遮挡样本。")
doc.add_picture(os.path.join(M,"chart_confusion.png"),width=Inches(4.6))

H("七、使用说明",1)
doc.add_paragraph("环境:Python 3.10,依赖 mindspore、mindcv、streamlit、pillow、openpyxl、python-pptx、matplotlib;",style="List Number")
doc.add_paragraph("数据准备:python prepare_data.py;",style="List Number")
doc.add_paragraph("训练:python train.py --epochs 30;",style="List Number")
doc.add_paragraph("推理:python predict.py <图像或目录>;",style="List Number")
doc.add_paragraph("前端:streamlit run app.py,浏览器访问 http://localhost:8501。",style="List Number")

H("八、局限与展望",1)
P("宏观分支的水色色块阈值为经验设定,后续需用本地海域实测海面照片校准;融合结论的置信度依赖问卷填写完整性,缺答时仅用图像证据、置信度相应降低。显微分支当前覆盖 7 个赤潮相关类群,后续可接入更多藻种与本地实测图像;识别目前基于单张图像,未来可结合细胞计数估算藻密度、给出赤潮发生概率;模型可进一步量化部署至昇腾端侧设备,服务现场监测。")

H("附录:数据溯源字段说明",1)
P("provenance.csv 逐张记录以下字段,确保每张训练图像可回溯到真实采集记录:slug(类别)、class_cn/class_sci(类别中文名/学名)、level(物种级/属级)、occurrence_key(GBIF 采集记录号)、record_scientificName(该记录实际学名)、media_url(影像原始 URL)、license(许可协议)、rightsHolder(权利人)、file(本地文件名)、bytes(字节数)。")

out=os.path.join(M,"设计文档.docx")
doc.save(out); print("saved",out)
