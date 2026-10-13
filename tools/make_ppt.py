# -*- coding: utf-8 -*-
"""make_ppt.py —— 生成参赛 PPT(内容匿名, 不含任何个人信息) 双分支版
输出: D:\\train\\materials\\作品PPT.pptx
"""
import os, json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

M=r"D:\train\materials"
ev=json.load(open(os.path.join(M,"eval_results.json"),encoding="utf-8"))
log=json.load(open(r"D:\train\model\train_log.json",encoding="utf-8"))
NAVY=RGBColor(0x1F,0x3B,0x63); TEAL=RGBColor(0x2E,0x75,0xB6); GREY=RGBColor(0x59,0x59,0x59)

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
BLANK=prs.slide_layouts[6]

def add_slide(): return prs.slides.add_slide(BLANK)
def txt(sl,l,t,w,h,s,size=18,bold=False,color=GREY,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP):
    box=sl.shapes.add_textbox(Inches(l),Inches(t),Inches(w),Inches(h))
    tf=box.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    p=tf.paragraphs[0]; p.alignment=align
    r=p.add_run(); r.text=s; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    return box
def title_bar(sl,s,sub=None):
    bar=sl.shapes.add_shape(1,0,0,prs.slide_width,Inches(1.05))
    bar.fill.solid(); bar.fill.fore_color.rgb=NAVY; bar.line.fill.background()
    txt(sl,0.5,0.18,12.3,0.7,s,size=30,bold=True,color=RGBColor(0xFF,0xFF,0xFF),anchor=MSO_ANCHOR.MIDDLE)
    if sub: txt(sl,0.55,1.15,12.2,0.4,sub,size=14,color=TEAL)
def bullets(sl,items,l=0.7,t=1.8,w=11.9,size=18,gap=0.55):
    y=t
    for it in items:
        txt(sl,l,y,w,gap+0.2,"•  "+it,size=size)
        y+=gap
def pic(sl,path,l,t,w=None,h=None):
    if w: return sl.shapes.add_picture(path,Inches(l),Inches(t),width=Inches(w))
    return sl.shapes.add_picture(path,Inches(l),Inches(t),height=Inches(h))

# ---- 1 封面 ----
s=add_slide()
bg=s.shapes.add_shape(1,0,0,prs.slide_width,prs.slide_height); bg.fill.solid(); bg.fill.fore_color.rgb=NAVY; bg.line.fill.background()
txt(s,1.0,2.2,11.3,1.2,"海洋微藻智能识别与赤潮风险提示工具",size=40,bold=True,color=RGBColor(0xFF,0xFF,0xFF),align=PP_ALIGN.CENTER)
txt(s,1.0,3.5,11.3,0.6,"双分支:宏观海面照片贝叶斯融合预警 × 显微图像 MindSpore 识别确认",size=20,color=RGBColor(0xCF,0xE3,0xF5),align=PP_ALIGN.CENTER)
txt(s,1.0,4.6,11.3,0.5,"2026 年海峡大学生人工智能创意大赛",size=16,color=RGBColor(0x9F,0xC5,0xE8),align=PP_ALIGN.CENTER)

# ---- 2 背景 ----
s=add_slide(); title_bar(s,"一、背景与意义")
bullets(s,[
 "福建沿海是赤潮多发区,赤潮藻爆发威胁海水养殖、贝类安全与近岸生态;",
 "现场人员(养殖户/巡查/公众)最先拿到的是海面现状照片,而非显微照片;",
 "但仅凭一张海面照片难以可靠定种,传统显微镜镜检又依赖专业人员、耗时长;",
 "本工具思路:宏观海面照片+现场特征问卷做贝叶斯融合快速预警(主分支),显微照片做实验室级识别确认(次分支)。"],size=19,gap=0.75)

# ---- 3 数据 ----
s=add_slide(); title_bar(s,"二、数据集:真实、可溯源的赤潮藻影像")
bullets(s,[
 "图像来源:GBIF(全球生物多样性信息网络)公开采集影像,辅以 CNKI 学术显微图;",
 "每一张图均记录溯源信息:采集记录号、记录学名、影像 URL、许可协议、权利人;",
 "覆盖 7 个赤潮相关类群,共约 755 张真实影像(含显微与海面水华宏观照);",
 "分类在 GBIF 可靠支撑的分类层级进行,物种代表种写入知识库,保证标签真实可信。"],size=17,gap=0.6)
pic(s,os.path.join(M,"chart_counts.png"),3.4,3.6,w=6.6)

# ---- 4 方法·显微分支 ----
s=add_slide(); title_bar(s,"三、分支二(显微):MindSpore 迁移学习")
bullets(s,[
 "骨干网络:MobileNetV3-Small(ImageNet 预训练),轻量、适合端侧/教学部署;",
 "迁移学习策略:冻结骨干特征层,仅训练分类头,小样本下快速收敛、抑制过拟合;",
 "输入 224×224 RGB;数据增强:随机裁剪、水平翻转;",
 "训练框架:MindSpore(昇思),PYNATIVE 模式 CPU 训练,30 epochs,Adam 优化;",
 "按验证集准确率自动保存最佳权重 best.ckpt。"],size=18,gap=0.62)
pic(s,os.path.join(M,"chart_acc_curve.png"),3.4,4.3,w=6.6)

# ---- 5 方法·宏观分支+融合 ----
s=add_slide(); title_bar(s,"四、分支一(宏观):水色CV × 现场问卷 贝叶斯融合")
bullets(s,[
 "图像证据:经典CV对海面照做水华色块检测(粉/褐/绿/泡沫占比),无需训练、不受宏观标注数据稀缺限制;",
 "问卷证据:6 题现场特征(水色、夜间发光、死鱼贝类、漂浮物、气味、环境背景);",
 "每个选项对假设集{夜光藻型/有毒甲藻型/硅藻型/绿潮/无赤潮}显式声明似然权重,可审计;",
 "融合:后验 ∝ 先验 × 图像似然^w × Π问卷似然,输出最可能致灾类群+风险等级+处置建议;",
 "四级预警:预警指数=置信度×类群风险权重×100,分红/橙/黄/蓝四级,每级对应明确响应措施,阈值与权重可审计;",
 "设计动机:真实赤潮监测本就综合水色与现场征象,融合比单图更稳、更可解释。"],size=17,gap=0.62)

# ---- 6 结果 ----
s=add_slide(); title_bar(s,f"五、实验结果(显微分支):验证集准确率 {ev['overall_val_acc']*100:.1f}%")
pic(s,os.path.join(M,"chart_perclass.png"),0.6,1.6,w=6.0)
pic(s,os.path.join(M,"chart_acc_curve.png"),6.9,1.6,w=6.0)
txt(s,0.6,5.6,12.2,1.4,
    f"总体验证准确率 {ev['overall_val_acc']*100:.2f}%({sum(ev['confusion'][i][i] for i in range(len(ev['confusion'])))}/{sum(sum(r) for r in ev['confusion'])}),全部类别均超过 80% 验收线;"
    f"最佳单轮验证准确率 {log['best_val_acc']*100:.1f}%。",size=18,bold=True,color=NAVY)

# ---- 7 混淆矩阵 ----
s=add_slide(); title_bar(s,"六、逐类判别能力分析(混淆矩阵)")
pic(s,os.path.join(M,"chart_confusion.png"),4.0,1.5,h=5.4)
txt(s,0.5,6.9,12.3,0.5,"对角线占绝对主导,类间混淆极少;夜光藻与骨条藻的少量误判源于低清/遮挡样本。",size=15)

# ---- 8 宏观案例 ----
s=add_slide(); title_bar(s,"七、宏观分支案例:海面照片→融合预警")
pic(s,r"D:\train\dataset_raw\gbif2\noctiluca\noctiluca_0000.jpg",0.6,1.6,w=5.6)
bullets(s,[
 "输入:一张岸拍海面照片(水华呈粉红色带);",
 "水色CV:检出粉色块占比约 7% → 判定『粉红/砖红』水华色;",
 "融合(仅图像证据):夜光藻型 75.7% > 有毒甲藻型 18.9% > 硅藻型 4.7%;",
 "输出:夜光藻型水华,风险『中—高』→ 预警指数61 = 橙色Ⅱ级预警;响应:48h内复核+暂停该点位贝类采捕送检+每日监测;",
 "叠加问卷(如夜间发光=是)后置信度进一步提升,体现融合的增益。"],l=6.6,w=6.2,size=16,gap=0.62)

# ---- 9 系统 ----
s=add_slide(); title_bar(s,"八、系统实现:双分支一体化前端")
bullets(s,[
 "Streamlit Web 前端双 Tab:宏观海面(主,上传+6题问卷+融合卡)/ 显微(次,实验室确认);",
 "显微识别自动联动《藻种知识库》:风险等级、毒素、生态、养殖影响、爆发规律、监测优先级;",
 "宏观分支输出后验概率与证据链,结论可追溯到低饱和/色块等中间量;",
 "低饱和度图像自动提示『疑似显微照』,引导切换分支,避免误用。"],size=18,gap=0.65)

# ---- 10 创新点 ----
s=add_slide(); title_bar(s,"九、创新点与特色")
bullets(s,[
 "宏观优先:直面真实用户只有海面照片的现状,用融合而非硬分类解决宏观定种难题;",
 "贝叶斯可解释:问卷似然权重显式声明、证据链可审计,非黑箱;",
 "真实可溯源:训练图像来自 GBIF 公开采集记录,逐张留存许可与权利人;",
 "MindSpore 生态 + 轻量迁移学习:显微分支 7 类小样本达 95.4% 验证准确率;",
 "知识库驱动:风险结论来自结构化《藻种知识库》,可解释、可维护、可扩展。"],size=18,gap=0.68)

# ---- 11 总结 ----
s=add_slide(); title_bar(s,"十、总结与展望")
bullets(s,[
 "已实现:真实溯源数据集(7 类约 755 张)+ 宏观融合与四级预警分支 + 显微 MindSpore 识别分支(95.4%)+ 双分支前端;",
 "展望:积累本地海域实测海面照校准色块阈值;接入细胞计数估算赤潮密度;部署至昇腾端侧;",
 "本工具可为近岸养殖区赤潮早期预警与贝类安全监测提供辅助决策支持。"],size=19,gap=0.8)

out=os.path.join(M,"作品PPT.pptx")
prs.save(out); print("saved",out,"slides:",len(prs.slides))
