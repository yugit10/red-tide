# -*- coding: utf-8 -*-
"""predict.py —— 加载 best.ckpt 对单张/批量图像做赤潮藻识别 + 风险提示
用法:
  python predict.py path/to/img.jpg
  python predict.py dir_of_images/
也可被 app.py 导入: from predict import Predictor
"""
import os, sys, json, glob
import numpy as np
from PIL import Image
import mindspore as ms
from mindspore import nn, context
from mindcv import create_model

DATA   = r"D:\train\algae_dataset"
MODEL  = r"D:\train\model\best.ckpt"
KB_XLSX= r"D:\train\藻种知识库.xlsx"
SIZE=224
MEAN=np.array([0.485,0.456,0.406],dtype=np.float32)
STD =np.array([0.229,0.224,0.225],dtype=np.float32)

def _softmax(x):
    x=x-x.max(axis=-1,keepdims=True)
    e=np.exp(x); return e/e.sum(axis=-1,keepdims=True)

def load_kb():
    """从《藻种知识库.xlsx》读取 slug->风险信息, 保证单一数据源"""
    kb={}
    try:
        from openpyxl import load_workbook
        wb=load_workbook(KB_XLSX,data_only=True); ws=wb.active
        head=[c.value for c in ws[1]]
        idx={h:i for i,h in enumerate(head)}
        for row in ws.iter_rows(min_row=2,values_only=True):
            slug_raw=row[idx.get("训练类别(slug)")]
            if not slug_raw: continue
            slug=str(slug_raw).split("(")[0].strip()
            kb[slug]={
              "cn":row[idx["藻种名称"]],"sci":row[idx["学名"]],
              "risk":row[idx["赤潮风险等级"]],"eco":row[idx["生态特性"]],
              "impact":row[idx["对养殖影响"]],"toxin":row[idx["毒素类型"]],
              "pattern":row[idx["爆发规律(季节/海域)"]],"priority":row[idx["监测优先级"]],
            }
    except Exception as e:
        print("[warn] 知识库加载失败:",e)
    return kb

class Predictor:
    def __init__(self):
        context.set_context(mode=context.PYNATIVE_MODE,device_target="CPU")
        cmap=json.load(open(os.path.join(DATA,"class_map.json"),encoding="utf-8"))
        self.classes=cmap["classes"]              # slug 顺序即标签0..N-1
        self.meta=cmap.get("meta",{})
        N=len(self.classes)
        net=create_model("mobilenet_v3_small_100",num_classes=N,pretrained=False,in_channels=3)
        if os.path.exists(MODEL):
            param=ms.load_checkpoint(MODEL)
            ms.load_param_into_net(net,param)
            print("[predictor] 已加载", MODEL)
        else:
            print("[predictor][warn] 未找到 best.ckpt, 使用未训练权重(结果无意义)")
        net.set_train(False); self.net=net
        self.kb=load_kb()
        self.best_acc=cmap.get("summary",{})

    def _prep(self,path):
        im=Image.open(path).convert("RGB")
        im=im.resize((SIZE,SIZE),Image.LANCZOS)
        a=np.asarray(im,dtype=np.float32)/255.0
        a=(a-MEAN)/STD
        a=a.transpose(2,0,1)[None,...]           # 1,3,224,224
        return ms.Tensor(a,ms.float32)

    def predict(self,path,topk=3):
        x=self._prep(path)
        logits=self.net(x).asnumpy()[0]
        probs=_softmax(logits)
        order=np.argsort(-probs)
        out=[]
        for i in order[:topk]:
            slug=self.classes[i]
            m=self.meta.get(slug,{})
            info=self.kb.get(slug,{})
            out.append({"slug":slug,"cn":info.get("cn") or m.get("cn",slug),
                        "sci":info.get("sci") or m.get("sci",""),
                        "prob":float(probs[i]),**{k:info.get(k) for k in
                        ("risk","eco","impact","toxin","pattern","priority")}})
        return out

def _main():
    if len(sys.argv)<2:
        print(__doc__); return
    p=Predictor()
    target=sys.argv[1]
    files=[target] if os.path.isfile(target) else sorted(glob.glob(os.path.join(target,"*")))
    files=[f for f in files if f.lower().endswith((".jpg",".jpeg",".png"))]
    for f in files[:20]:
        r=p.predict(f,topk=3)
        top=r[0]
        print(f"\n{os.path.basename(f)}")
        for j,x in enumerate(r):
            mark=">>" if j==0 else "  "
            print(f"{mark} {x['cn']}({x['slug']}) {x['prob']*100:5.1f}%  风险:{x.get('risk','?')}")
        print(f"   毒素:{top.get('toxin')} | 监测:{top.get('priority')}")

if __name__=="__main__":
    _main()
