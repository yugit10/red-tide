# -*- coding: utf-8 -*-
import docx, sys

# 设计文档标题结构
d = docx.Document(r'D:\train\materials\设计文档.docx')
out = ['===== 设计文档段落结构 =====']
for p in d.paragraphs:
    t = p.text.strip()
    if not t:
        continue
    st = p.style.name if p.style else ''
    if st.lower().startswith('heading') or st.lower().startswith('标题'):
        lvl = st.lower().replace('heading', '').replace('标题', '')
        out.append('[H%s] %s' % (lvl, t))
    elif t.startswith(('一、', '二、', '三、', '四、', '五、', '六、', '七、', '八、', '九、', '十、',
                       '（一）', '（二）', '（三）', '（四）', '（五）', '（六）',
                       '(一)', '(二)', '(三)', '(四)', '(五)', '(六)')) and len(t) < 40:
        out.append('[节] ' + t)
out.append('')
out.append('===== 关键词检查 =====')
full = '\n'.join(p.text for p in d.paragraphs if p.text)
for kw in ['接口', '数据结构', '系统架构', '模块', '加分', 'MindSpore', '昇思', '昇腾', 'provenance', '溯源']:
    out.append('%s: %d 次' % (kw, full.count(kw)))
out.append('')
out.append('===== 表格数量 =====')
out.append('tables: %d' % len(d.tables))

# PPT 结构
from pptx import Presentation
prs = Presentation(r'D:\train\materials\作品PPT.pptx')
out.append('')
out.append('===== PPT 页数: %d =====' % len(prs.slides))
for i, slide in enumerate(prs.slides, 1):
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip().replace('\n', ' / ')
            if t:
                texts.append(t[:60])
    title = texts[0] if texts else '(无文本)'
    out.append('第%d页: %s' % (i, title))
    for t in texts[1:3]:
        out.append('    - %s' % t)

with open(r'D:\train\_audit_structure.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print('DONE')
